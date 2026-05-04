#!/usr/bin/env python3
"""Read 'Chart Tank Data.xlsx' and regenerate chart-data.json.

v2: also reads Hero Image Path, VJV-Right/Left Part # (Tanks sheet) and
Image Path (Add-On Options sheet).
"""
import json
from pathlib import Path
from openpyxl import load_workbook

REPO = Path(__file__).resolve().parent
XLSX = REPO / "Chart Tank Data.xlsx"
OUT = REPO / "chart-data.json"


def truthy(v) -> bool:
    if v is None:
        return False
    if isinstance(v, bool):
        return v
    s = str(v).strip().lower()
    return s in ("true", "yes", "y", "1", "x")


def s(v) -> str:
    return "" if v is None else str(v).strip()


def rows(ws, header_row=1):
    headers = [s(c.value) for c in ws[header_row]]
    for r in range(header_row + 1, ws.max_row + 1):
        cells = [ws.cell(row=r, column=ci).value for ci in range(1, len(headers) + 1)]
        if all(c is None or s(c) == "" for c in cells):
            continue
        yield dict(zip(headers, cells))


def main() -> None:
    wb = load_workbook(XLSX, data_only=True)

    # Tanks
    tanks: dict[str, dict] = {}
    for r in rows(wb["Tanks"]):
        tid = s(r["Tank ID"])
        if not tid:
            continue
        t = {
            "id": tid,
            "displayName": s(r["Display Name"]),
            "size": s(r["Size"]),
            "pressureClass": s(r["Pressure Class"]),
            "fillType": s(r["Fill Type"]),
            "psiRating": s(r["PSI Rating"]),
            "standard": s(r["Standard (DOT/ASME)"]),
            "sourceFile": s(r["Source PDF"]),
            "heroImage": s(r.get("Hero Image Path", "")),
            "configSteps": [],
            "standardIncludes": [],
            "applicableOptions": [],
        }
        vjvr = s(r.get("VJV-Right Part #", ""))
        vjvl = s(r.get("VJV-Left Part #", ""))
        if vjvr:
            t["vjvRight"] = {"partNumber": vjvr}
        if vjvl:
            t["vjvLeft"] = {"partNumber": vjvl}
        tanks[tid] = t

    # Config steps
    step_buckets: dict[str, dict[int, dict]] = {tid: {} for tid in tanks}
    for r in rows(wb["Config Steps"]):
        tid = s(r["Tank ID"])
        if tid not in tanks:
            continue
        step_order = int(r["Step Order"] or 0)
        opt_order = int(r["Option Order"] or 0)
        bucket = step_buckets[tid].setdefault(step_order, {
            "stepName": s(r["Step Name"]),
            "selectionType": s(r["Selection Type"]) or "single",
            "_options": [],
        })
        bucket["_options"].append((opt_order, {
            "label": s(r["Option Label"]),
            "partNumber": s(r["Part Number"]),
            "notes": s(r["Notes"]),
        }))
    for tid, t in tanks.items():
        for so in sorted(step_buckets[tid]):
            step = step_buckets[tid][so]
            opts = [o for _, o in sorted(step["_options"])]
            t["configSteps"].append({
                "stepName": step["stepName"],
                "selectionType": step["selectionType"],
                "options": opts,
            })

    # Standard Includes
    for r in rows(wb["Standard Includes"]):
        tid = s(r["Tank ID"])
        if tid not in tanks:
            continue
        item = {
            "label": s(r["Item Label"]),
            "partNumber": s(r["Part Number"]),
            "notes": s(r["Notes"]),
        }
        src = r.get("Source Option #")
        if src not in (None, ""):
            try:
                item["sourceOption"] = int(src)
            except Exception:
                pass
        tanks[tid]["standardIncludes"].append(item)

    # Applicability
    for r in rows(wb["Applicability"]):
        tid = s(r["Tank ID"])
        if tid not in tanks:
            continue
        try:
            n = int(r["Option #"])
        except Exception:
            continue
        tanks[tid]["applicableOptions"].append({
            "optionNumber": n,
            "applies": truthy(r["Applies"]),
            "standard": truthy(r["Standard (Auto-Include)"]),
            "notes": s(r["Notes"]),
        })
    for t in tanks.values():
        t["applicableOptions"].sort(key=lambda x: x["optionNumber"])

    # Add-On Options + Variants
    options: dict[int, dict] = {}
    for r in rows(wb["Add-On Options"]):
        try:
            n = int(r["Option #"])
        except Exception:
            continue
        options[n] = {
            "number": n,
            "name": s(r["Name"]),
            "shortDescription": s(r["Short Description"]),
            "selectionType": s(r["Selection Type"]) or "yes-no",
            "image": s(r.get("Image Path", "")),
            "variants": [],
        }
    var_buckets: dict[int, list] = {n: [] for n in options}
    for r in rows(wb["Add-On Variants"]):
        try:
            n = int(r["Option #"])
        except Exception:
            continue
        if s(r["Variant Label"]).startswith("(no variants)"):
            continue
        try:
            order = int(r["Variant Order"] or 0)
        except Exception:
            order = 0
        var_buckets.setdefault(n, []).append((order, {
            "label": s(r["Variant Label"]),
            "partNumber": s(r["Part Number"]),
            "notes": s(r["Notes"]),
        }))
    for n, opt in options.items():
        opt["variants"] = [v for _, v in sorted(var_buckets.get(n, []))]
        for v in opt["variants"]:
            if not v.get("notes"):
                v.pop("notes", None)

    out = {
        "_meta": {
            "version": "2.0",
            "generated_by": "xlsx_to_json.py",
            "tankCount": len(tanks),
            "issues": [],
        },
        "tanks": list(tanks.values()),
        "addOnOptions": [options[k] for k in sorted(options)],
    }
    OUT.write_text(json.dumps(out, indent=2))
    print(f"Wrote {OUT.name}")
    print(f"  tanks: {len(tanks)}  add-ons: {len(options)}")


if __name__ == "__main__":
    main()
