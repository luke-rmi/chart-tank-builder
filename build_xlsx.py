#!/usr/bin/env python3
"""Rebuild 'Chart Tank Data.xlsx' from chart-data.json.

Run this after a schema change. Day-to-day editing flows the OTHER way:
edit the .xlsx, then run xlsx_to_json.py to regenerate chart-data.json.

v2 schema additions:
  • Tanks sheet — Hero Image Path, VJV-Right Part #, VJV-Left Part # columns.
  • Add-On Options sheet — Image Path column.
"""
import json
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

REPO = Path(__file__).resolve().parent
DATA = json.loads((REPO / "chart-data.json").read_text())
OUT = REPO / "Chart Tank Data.xlsx"

HEADER_FILL = PatternFill("solid", start_color="0A4D8C")
HEADER_FONT = Font(name="Arial", color="FFFFFF", bold=True, size=11)
BODY_FONT = Font(name="Arial", size=10)
TITLE_FONT = Font(name="Arial", size=16, bold=True, color="073968")
SUBTITLE_FONT = Font(name="Arial", size=11, italic=True, color="4A5568")
NOTE_FONT = Font(name="Arial", size=10, color="4A5568", italic=True)
THIN = Side(border_style="thin", color="DBDBDB")
CELL_BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
ALT_FILL = PatternFill("solid", start_color="F7F9FC")
STD_FILL = PatternFill("solid", start_color="EEF7F0")
STD_FONT = Font(name="Arial", size=10, color="1A7F3C", bold=True)


def write_header(ws, row, headers, widths=None):
    for ci, h in enumerate(headers, 1):
        c = ws.cell(row=row, column=ci, value=h)
        c.font = HEADER_FONT
        c.fill = HEADER_FILL
        c.alignment = Alignment(horizontal="left", vertical="center")
        c.border = CELL_BORDER
    ws.row_dimensions[row].height = 24
    if widths:
        for ci, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(ci)].width = w


def write_row(ws, row, values, alt=False, style_func=None):
    for ci, v in enumerate(values, 1):
        c = ws.cell(row=row, column=ci, value=v)
        c.font = BODY_FONT
        c.alignment = Alignment(vertical="top", wrap_text=True)
        c.border = CELL_BORDER
        if alt:
            c.fill = ALT_FILL
        if style_func:
            style_func(ci, c, v)


wb = Workbook()

# README
ws = wb.active
ws.title = "README"
ws.column_dimensions["A"].width = 110
ws["A1"] = "Chart Microbulk Tank Builder · Data Workbook"
ws["A1"].font = TITLE_FONT
ws["A2"] = (f"Generated from chart-data.json · v{DATA['_meta'].get('version','?')} · "
            f"{DATA['_meta'].get('tankCount', '?')} tanks · "
            f"{len(DATA['addOnOptions'])} add-on options")
ws["A2"].font = SUBTITLE_FONT
notes = [
    "",
    "Day-to-day edit flow:",
    "  1. Edit values on the data sheets.",
    "  2. Save the file.",
    "  3. Run ./regenerate.sh — it converts the xlsx back to chart-data.json.",
    "  4. Commit + push from Cursor — GitHub Pages serves the new data within ~60s.",
    "",
    "Sheet guide:",
    "  • Tanks – the master tank list. Each row = one configurator entry. id and displayName must be unique.",
    "          New v2 columns: Hero Image Path, VJV-Right Part #, VJV-Left Part #.",
    "  • Config Steps – every selectable option per tank, in display order. Step Order controls the sequence.",
    "          (Tip) The configurator auto-skips any step that has only ONE option, so 1-row Tank Model entries no longer cause a click.",
    "  • Add-On Options – the 18 catalog options shown after the per-tank steps. New v2 column: Image Path.",
    "  • Add-On Variants – the part-number sub-options for each add-on (gas variants, top/flex variants, etc.).",
    "  • Applicability – per-tank rules for each add-on (Applies / Standard / Notes).",
    "  • Standard Includes – auto-added items shown in the summary's 'Standard (Included)' section.",
    "",
    "Conventions:",
    "  • TRUE / FALSE in the boolean columns. Anything else is treated as FALSE.",
    "  • Part numbers are case-sensitive. Don't add stray spaces.",
    "  • Image paths are relative to the repo root, e.g. images/tanks/5500L-MP-FF.webp.",
    "  • Tank IDs follow <size>-<class>-<fill>, e.g. 5500L-MP-FF, 2000L-HP-TF, permamax-2200lb-HP-CO2.",
    "",
    "VJV (Vacuum-Jacketed Valve) handling — v2 change:",
    "  • Five tanks (1000/1500/2000/3000 HP TopFill + 5500L MP FlexFill) USED to expose VJ-L/VJ-R rows in the Tank Model step.",
    "  • These have been removed from Tank Model. The configurator now reads VJV-Right Part # and VJV-Left Part #",
    "    from the Tanks sheet and feeds them dynamically into Option #12 when the tank is selected.",
    "",
    "Questions: webtech@rmimfg.com.",
]
for i, line in enumerate(notes, 4):
    cell = ws.cell(row=i, column=1, value=line)
    if line.startswith(("Day-to-day", "Sheet guide", "Conventions", "VJV", "Questions")):
        cell.font = Font(name="Arial", size=12, bold=True, color="073968")
    else:
        cell.font = NOTE_FONT
ws.sheet_view.showGridLines = False


# Tanks
ws = wb.create_sheet("Tanks")
write_header(ws, 1, [
    "Tank ID", "Display Name", "Size", "Pressure Class", "Fill Type",
    "PSI Rating", "Standard (DOT/ASME)", "Source PDF",
    "Hero Image Path", "VJV-Right Part #", "VJV-Left Part #"
], widths=[26, 38, 12, 14, 12, 12, 14, 36, 36, 28, 28])
ws.freeze_panes = "A2"
for i, t in enumerate(DATA["tanks"], 2):
    write_row(ws, i, [
        t["id"], t["displayName"], t["size"], t["pressureClass"],
        t["fillType"], t["psiRating"], t.get("standard", ""), t["sourceFile"],
        t.get("heroImage", ""),
        t.get("vjvRight", {}).get("partNumber", "") if isinstance(t.get("vjvRight"), dict) else "",
        t.get("vjvLeft", {}).get("partNumber", "") if isinstance(t.get("vjvLeft"), dict) else "",
    ], alt=(i % 2 == 0))


# Config Steps
ws = wb.create_sheet("Config Steps")
write_header(ws, 1, [
    "Tank ID", "Step Order", "Step Name", "Selection Type",
    "Option Order", "Option Label", "Part Number", "Notes"
], widths=[26, 10, 44, 14, 12, 60, 36, 24])
ws.freeze_panes = "A2"
row = 2
for t in DATA["tanks"]:
    for si, step in enumerate(t["configSteps"], 1):
        for oi, opt in enumerate(step["options"], 1):
            write_row(ws, row, [
                t["id"], si, step["stepName"], step.get("selectionType", "single"),
                oi, opt["label"], opt["partNumber"], opt.get("notes", ""),
            ], alt=(row % 2 == 0))
            row += 1


# Add-On Options
ws = wb.create_sheet("Add-On Options")
write_header(ws, 1, [
    "Option #", "Name", "Short Description", "Selection Type", "Image Path"
], widths=[10, 38, 60, 22, 36])
ws.freeze_panes = "A2"
for i, opt in enumerate(DATA["addOnOptions"], 2):
    write_row(ws, i, [
        opt["number"], opt["name"], opt["shortDescription"],
        opt["selectionType"], opt.get("image", ""),
    ], alt=(i % 2 == 0))


# Add-On Variants
ws = wb.create_sheet("Add-On Variants")
write_header(ws, 1, [
    "Option #", "Variant Order", "Variant Label", "Part Number", "Notes"
], widths=[10, 12, 30, 38, 40])
ws.freeze_panes = "A2"
row = 2
for opt in DATA["addOnOptions"]:
    for i, v in enumerate(opt.get("variants", []), 1):
        write_row(ws, row, [
            opt["number"], i, v["label"], v["partNumber"], v.get("notes", ""),
        ], alt=(row % 2 == 0))
        row += 1
    if not opt.get("variants"):
        write_row(ws, row, [opt["number"], 0, "(no variants)", "", "Yes/No only"], alt=(row % 2 == 0))
        row += 1


# Applicability
ws = wb.create_sheet("Applicability")
write_header(ws, 1, [
    "Tank ID", "Option #", "Option Name", "Applies", "Standard (Auto-Include)", "Notes"
], widths=[26, 10, 38, 10, 22, 50])
ws.freeze_panes = "A2"
row = 2
opt_name = {o["number"]: o["name"] for o in DATA["addOnOptions"]}
for t in DATA["tanks"]:
    for ap in t["applicableOptions"]:
        std = ap["standard"]
        def style(ci, c, v, std=std):
            if std and ci in (4, 5):
                c.fill = STD_FILL
                c.font = STD_FONT
        write_row(ws, row, [
            t["id"], ap["optionNumber"], opt_name.get(ap["optionNumber"], ""),
            "TRUE" if ap["applies"] else "FALSE",
            "TRUE" if std else "FALSE",
            ap.get("notes", ""),
        ], alt=(row % 2 == 0), style_func=style)
        row += 1


# Standard Includes
ws = wb.create_sheet("Standard Includes")
write_header(ws, 1, [
    "Tank ID", "Item Label", "Part Number", "Source Option #", "Notes"
], widths=[26, 44, 36, 14, 40])
ws.freeze_panes = "A2"
row = 2
for t in DATA["tanks"]:
    for s in t.get("standardIncludes", []):
        write_row(ws, row, [
            t["id"], s["label"], s.get("partNumber", ""),
            s.get("sourceOption", ""), s.get("notes", ""),
        ], alt=(row % 2 == 0))
        row += 1

wb.move_sheet("README", offset=-wb.sheetnames.index("README"))
wb.save(OUT)
print(f"✓ Wrote {OUT.name}")
print(f"  Sheets: {', '.join(wb.sheetnames)}")
